import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation

from app.core.config import settings
from app.main import app
from app.models import CourseEnrollment, Material
from app.models.user import UserRole
from app.services.storage import LocalStorage, get_storage

API = "/api/v1"


def pdf_bytes(tag: str = "a") -> bytes:
    return f"%PDF-1.4\n% {tag}\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n".encode()


def docx_bytes(text: str = "hello") -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes(text: str = "slide") -> bytes:
    buffer = io.BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = text
    presentation.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def storage(tmp_path):
    store = LocalStorage(tmp_path)
    app.dependency_overrides[get_storage] = lambda: store
    yield store
    app.dependency_overrides.pop(get_storage, None)


def stored_files(storage: LocalStorage) -> list:
    return [p for p in storage.root.rglob("*") if p.is_file()]


@pytest.fixture
def teacher(make_user):
    return make_user(UserRole.TEACHER)


@pytest.fixture
def course_and_topic(client, teacher, auth_headers):
    h = auth_headers(teacher)
    course = client.post(f"{API}/courses", json={"code": "CS301", "title": "DB"}, headers=h).json()
    topic = client.post(
        f"{API}/courses/{course['id']}/topics", json={"title": "SQL Joins"}, headers=h
    ).json()
    return course, topic


def upload(client, headers, course_id, topic_id, name, content, mime="application/octet-stream"):
    return client.post(
        f"{API}/courses/{course_id}/topics/{topic_id}/materials",
        files={"file": (name, content, mime)},
        headers=headers,
    )


@pytest.mark.parametrize(
    ("name", "content_factory", "file_type"),
    [
        ("lecture.pdf", pdf_bytes, "pdf"),
        ("notes.docx", docx_bytes, "docx"),
        ("deck.PPTX", pptx_bytes, "pptx"),
    ],
)
def test_valid_uploads(
    client, storage, teacher, course_and_topic, auth_headers, name, content_factory, file_type
):
    course, topic = course_and_topic
    content = content_factory()
    r = upload(client, auth_headers(teacher), course["id"], topic["id"], name, content)
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["file_type"] == file_type
    assert body["status"] == "uploaded"
    assert body["file_size_bytes"] == len(content)
    assert body["uploaded_by"] == str(teacher.id)
    assert "storage_path" not in body
    assert len(stored_files(storage)) == 1


def test_download_returns_same_bytes(client, storage, teacher, course_and_topic, auth_headers):
    course, topic = course_and_topic
    h = auth_headers(teacher)
    content = pdf_bytes("dl")
    material = upload(client, h, course["id"], topic["id"], "my file.pdf", content).json()
    r = client.get(f"{API}/materials/{material['id']}/download", headers=h)
    assert r.status_code == 200
    assert r.content == content
    assert r.headers["content-type"] == "application/pdf"
    assert "my%20file.pdf" in r.headers["content-disposition"]


def test_rejects_wrong_extension(client, storage, teacher, course_and_topic, auth_headers):
    course, topic = course_and_topic
    for name in ("virus.exe", "noext", "notes.txt"):
        r = upload(client, auth_headers(teacher), course["id"], topic["id"], name, b"data")
        assert r.status_code == 422, name
    assert stored_files(storage) == []


def test_rejects_spoofed_content(client, storage, teacher, course_and_topic, auth_headers):
    course, topic = course_and_topic
    h = auth_headers(teacher)
    cases = [
        ("fake.pdf", b"MZ not a pdf"),
        ("fake.docx", pdf_bytes()),
        ("fake.docx", pptx_bytes()),
        ("fake.pptx", docx_bytes()),
        ("broken.docx", b"PK\x03\x04 not really a zip"),
    ]
    for name, content in cases:
        r = upload(client, h, course["id"], topic["id"], name, content)
        assert r.status_code == 422, name
    assert stored_files(storage) == []


def test_rejects_empty_file(client, storage, teacher, course_and_topic, auth_headers):
    course, topic = course_and_topic
    r = upload(client, auth_headers(teacher), course["id"], topic["id"], "empty.pdf", b"")
    assert r.status_code == 422
    assert stored_files(storage) == []


def test_rejects_oversize_file(
    client, storage, teacher, course_and_topic, auth_headers, monkeypatch
):
    monkeypatch.setattr(settings, "max_upload_mb", 1)
    course, topic = course_and_topic
    big = pdf_bytes() + b"0" * (1024 * 1024)
    r = upload(client, auth_headers(teacher), course["id"], topic["id"], "big.pdf", big)
    assert r.status_code == 422 and "limit" in r.json()["detail"]
    assert stored_files(storage) == []


def test_duplicate_file_conflicts_and_leaves_no_extra_file(
    client, storage, teacher, course_and_topic, auth_headers
):
    course, topic = course_and_topic
    h = auth_headers(teacher)
    content = pdf_bytes("same")
    first = upload(client, h, course["id"], topic["id"], "a.pdf", content).json()
    r = upload(client, h, course["id"], topic["id"], "renamed.pdf", content)
    assert r.status_code == 409
    assert r.json()["existing_material_id"] == first["id"]
    assert len(stored_files(storage)) == 1


def test_same_file_allowed_after_delete(client, storage, teacher, course_and_topic, auth_headers):
    course, topic = course_and_topic
    h = auth_headers(teacher)
    content = pdf_bytes("again")
    first = upload(client, h, course["id"], topic["id"], "a.pdf", content).json()
    assert client.delete(f"{API}/materials/{first['id']}", headers=h).status_code == 204
    assert upload(client, h, course["id"], topic["id"], "a.pdf", content).status_code == 201


def test_topic_must_belong_to_course(client, storage, teacher, course_and_topic, auth_headers):
    course, _ = course_and_topic
    h = auth_headers(teacher)
    other = client.post(f"{API}/courses", json={"code": "C2", "title": "Other"}, headers=h).json()
    foreign_topic = client.post(
        f"{API}/courses/{other['id']}/topics", json={"title": "X"}, headers=h
    ).json()
    r = upload(client, h, course["id"], foreign_topic["id"], "a.pdf", pdf_bytes())
    assert r.status_code == 404
    assert stored_files(storage) == []


def test_permissions(client, storage, teacher, make_user, course_and_topic, auth_headers, db):
    course, topic = course_and_topic
    material = upload(
        client, auth_headers(teacher), course["id"], topic["id"], "a.pdf", pdf_bytes()
    ).json()

    stranger = auth_headers(make_user(UserRole.TEACHER))
    assert (
        upload(client, stranger, course["id"], topic["id"], "b.pdf", pdf_bytes("b")).status_code
        == 404
    )
    assert (
        client.get(f"{API}/courses/{course['id']}/materials", headers=stranger).status_code == 404
    )
    assert (
        client.get(f"{API}/materials/{material['id']}/download", headers=stranger).status_code
        == 404
    )
    assert client.delete(f"{API}/materials/{material['id']}", headers=stranger).status_code == 404

    student = make_user(UserRole.STUDENT)
    db.add(CourseEnrollment(course_id=course["id"], student_id=student.id))
    db.commit()
    sh = auth_headers(student)
    assert upload(client, sh, course["id"], topic["id"], "c.pdf", pdf_bytes("c")).status_code == 403
    assert client.get(f"{API}/courses/{course['id']}/materials", headers=sh).status_code == 403
    assert client.get(f"{API}/materials/{material['id']}/download", headers=sh).status_code == 403

    admin = auth_headers(make_user(UserRole.ADMIN))
    assert (
        client.get(f"{API}/materials/{material['id']}/download", headers=admin).status_code == 200
    )

    assert client.get(f"{API}/courses/{course['id']}/materials").status_code == 401


def test_list_filters_by_topic_and_hides_deleted(
    client, storage, teacher, course_and_topic, auth_headers
):
    course, topic = course_and_topic
    h = auth_headers(teacher)
    topic2 = client.post(
        f"{API}/courses/{course['id']}/topics", json={"title": "Indexes"}, headers=h
    ).json()
    m1 = upload(client, h, course["id"], topic["id"], "a.pdf", pdf_bytes("1")).json()
    m2 = upload(client, h, course["id"], topic2["id"], "b.pdf", pdf_bytes("2")).json()

    all_ids = {
        m["id"] for m in client.get(f"{API}/courses/{course['id']}/materials", headers=h).json()
    }
    assert all_ids == {m1["id"], m2["id"]}
    by_topic = client.get(
        f"{API}/courses/{course['id']}/materials?topic_id={topic['id']}", headers=h
    ).json()
    assert [m["id"] for m in by_topic] == [m1["id"]]

    client.delete(f"{API}/materials/{m1['id']}", headers=h)
    left = client.get(f"{API}/courses/{course['id']}/materials", headers=h).json()
    assert [m["id"] for m in left] == [m2["id"]]
    assert client.get(f"{API}/materials/{m1['id']}/download", headers=h).status_code == 404


def test_deleting_topic_hides_its_materials(
    client, storage, teacher, course_and_topic, auth_headers
):
    course, topic = course_and_topic
    h = auth_headers(teacher)
    upload(client, h, course["id"], topic["id"], "a.pdf", pdf_bytes())
    client.delete(f"{API}/topics/{topic['id']}", headers=h)
    assert client.get(f"{API}/courses/{course['id']}/materials", headers=h).json() == []


def test_file_name_is_sanitized(client, storage, teacher, course_and_topic, auth_headers):
    course, topic = course_and_topic
    r = upload(
        client,
        auth_headers(teacher),
        course["id"],
        topic["id"],
        "../../etc/passwd.pdf",
        pdf_bytes(),
    )
    assert r.status_code == 201
    assert r.json()["file_name"] == "passwd.pdf"
    assert all(p.is_relative_to(storage.root) for p in stored_files(storage))


def test_stored_file_removed_when_database_write_fails(
    client, db, storage, teacher, course_and_topic, auth_headers, monkeypatch
):
    course, topic = course_and_topic

    def boom():
        raise RuntimeError("db down")

    monkeypatch.setattr(db, "commit", boom)
    # `client` fixture installs the get_db override; reuse it with a non-raising client
    client = TestClient(app, raise_server_exceptions=False)
    r = upload(client, auth_headers(teacher), course["id"], topic["id"], "a.pdf", pdf_bytes())
    assert r.status_code == 500
    assert stored_files(storage) == []
    assert db.query(Material).count() == 0


def test_local_storage_rejects_path_traversal(tmp_path):
    store = LocalStorage(tmp_path)
    with pytest.raises(ValueError):
        store.save("../escape.pdf", [b"x"])
    assert not (tmp_path.parent / "escape.pdf").exists()
